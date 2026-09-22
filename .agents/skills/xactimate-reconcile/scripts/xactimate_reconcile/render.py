"""One report supplies every amount in Markdown, DOCX and PDF drafts."""
from __future__ import annotations
from pathlib import Path
import html
import re
from .store import within, sha_file


def p(text):
    return {'type': 'p', 'text': str(text)}


def heading(text):
    return {'type': 'heading', 'text': str(text)}


def table(headers, rows):
    return {'type': 'table', 'headers': headers, 'rows': [[str(x) if x is not None else 'Unknown' for x in row] for row in rows]}


def dollars(value):
    if value is None:
        return 'Unknown'
    from decimal import Decimal
    return f'${Decimal(value):,.2f}'


def build_documents(case, report, context=None):
    if report.get('status') == 'blocked':
        blocks = [p('BLOCKED COMPARISON — DO NOT SUBMIT'),p('Resolve the errors below and regenerate. No revision amount is advanced from this run.')]
        blocks += [p(str(i.get('ref','case')) + ': ' + str(i.get('message','Review input')))
                   for i in report.get('issues',[]) if isinstance(i,dict)]
        return {key:{'title':title,'blocks':list(blocks)} for key,title in (
            ('reconciliation_memo','Blocked Reconciliation Review'),('cover_letter','Cover Letter Withheld'),
            ('operator_checklist','Input Correction Checklist'),('evidence_appendix','Source Validation Findings'))}
    cid = case.get('case_id', 'Unidentified case')
    state = 'SYNTHETIC EXAMPLE - NOT FOR SUBMISSION' if case.get('synthetic') else 'DRAFT FOR REVIEW'
    totals = report.get('totals')
    gs = {g['claim_id']: g for g in case.get('groups', [])}
    computed = [g for g in report.get('groups', []) if g.get('delta') is not None]
    evidence = {e['id']: e for e in case.get('evidence', [])}
    docs = {d['id']: d for d in case.get('documents', [])}
    basis = 'direct construction costs before taxes and general overhead and profit'
    if totals:
        opening = f"Please review the documented changes from {dollars(totals['baseline_direct'])} to {dollars(totals['proposed_direct'])} in {basis}, a net change of {dollars(totals['delta_direct'])}."
    else:
        opening = 'Please review the identified estimate differences and the outstanding evidence requests. The comparison is incomplete; no complete revision amount or payment request is stated.'
    memo = [p(state), p(f"Case {cid} | Comparison date {case.get('as_of','Unconfirmed')} | {case.get('author_role','Technical preparation')}"), p(opening),
            p('The figures concern the specified construction work. They do not determine coverage or the amount currently payable.'), heading('Estimates and record')]
    memo.append(table(['Role', 'Document', 'Version', 'Recorded date'], [[d['role'], d['id'], d.get('version'), d.get('date')] for d in docs.values() if d['role'] in ('baseline','proposed')]))
    memo.append(p('Confirm the date printed on each estimate. Intake receipt dates alone do not establish its pricing date.'))
    memo.append(p(f"Pricing context: {case.get('price_context',{}).get('location','Unconfirmed')}; list {case.get('price_context',{}).get('list') or 'not supplied'}; month {case.get('price_context',{}).get('month') or 'not supplied'}."))
    memo += [heading('Discrepancy schedule'), table(['Issue', 'Existing', 'Proposed', 'Change'], [[g['claim_id'], dollars(g.get('baseline_direct')), dollars(g.get('proposed_direct')), dollars(g.get('delta'))] for g in report.get('groups', [])])]
    for item in report.get('groups', []):
        raw = gs.get(item['claim_id'], {})
        memo.append(heading(item['claim_id'] + ' ' + raw.get('title', raw.get('roof_location','Roof work')).replace('_',' ')))
        memo.append(p(raw.get('observation') or raw.get('mapping_rationale') or 'Observation still required.'))
        memo.append(p(raw.get('necessity') or 'Confirm the physical basis for the proposed operation.'))
        memo.append(p('Comparison: ' + str(item.get('mapping_rationale', raw.get('mapping_rationale','Unresolved')))))
        if item.get('delta') is not None:
            memo.append(p(f"Calculation: {dollars(item.get('proposed_direct'))} less {dollars(item.get('baseline_direct'))} = {dollars(item['delta'])}."))
            effects = [f"{k.replace('_',' ')} {dollars(v)}" for k,v in item.get('components',{}).items() if v not in ('0.00',None)]
            if effects:
                memo.append(p('Attribution: ' + '; '.join(effects) + '.'))
        else:
            memo.append(p('Amount held pending resolution of the identified inputs.'))
        memo.append(p('Evidence: ' + ', '.join(item.get('source_ids', raw.get('source_ids',[]))) + '.'))
        memo.append(p('Competing explanation: ' + str(item.get('counterargument', raw.get('counterargument','Not supplied.')))))
        memo.append(p('Response or concession: ' + str(raw.get('reply') or raw.get('disposition') or 'Obtain the missing evidence before advancing this item.')))
        memo.append(p('Requested action: ' + str(raw.get('requested_action') or 'Provide an itemized response or identify the evidence required.')))
        memo.append(p('Scope: ' + str(item.get('scope_status',raw.get('scope_support',{}).get('status','unknown'))) + '. Coverage: ' + str(item.get('coverage_status',raw.get('coverage',{}).get('status','unknown'))) + '.'))
    memo += [heading('Financial reconciliation'), p('The direct-cost comparison and subsequent estimate adjustments are separate from policy settlement.')]
    bridge = report.get('financial_bridge', {})
    for side in ('baseline','proposed'):
        detail = bridge.get(side) or {}
        memo.append(p(side.capitalize() + ': direct costs ' + dollars(detail.get('direct')) + '; estimate total on the specified adjustment basis ' + dollars(detail.get('total')) + '. Basis: ' + str(detail.get('basis_status','not supplied')).replace('_',' ') + '.'))
        for adjustment in detail.get('adjustments',[]):
            label = str(adjustment.get('label') or adjustment.get('id'))
            explanation = ('fraction ' + str(adjustment.get('input')) + ' applied to ' + dollars(adjustment.get('basis_total'))) if adjustment.get('method')=='rate' else ('supplied amount ' + str(adjustment.get('input')))
            memo.append(p(label + ': ' + explanation + ' = ' + dollars(adjustment.get('calculated')) + '. Basis items: ' + ', '.join(adjustment.get('basis') or []) + '.'))
    scenario = bridge.get('payment_scenario')
    if scenario:
        memo.append(p('Supplied payment scenario: ' + str(scenario.get('label','Scenario')) + '; status ' + str(scenario.get('status','held')) + '; result ' + dollars(scenario.get('result')) + '.'))
        for step in scenario.get('steps',[]):
            memo.append(p('Scenario step ' + str(step.get('id')) + ': ' + str(step.get('operation')) + ' ' + dollars(step.get('amount')) + '; resulting balance ' + dollars(step.get('after')) + '.'))
    else:
        memo.append(p('Payment scenario: not supplied or not calculable.'))
    memo.append(p('Amount currently payable: not determined.'))
    memo.append(p('Absent tax or markup inputs mean not specified, rather than verified zero. Any payment scenario is an explicit calculation supplied for review; it is not a coverage decision.'))
    memo += [heading('Open items and next action')]
    for issue in report.get('issues', []):
        memo.append(p(f"{issue.get('ref','case')}: {issue.get('message','Review required')}"))
    if not report.get('issues'):
        memo.append(p('Confirm source readings, sender authority, and the licensed Xactimate entry and export before release.'))
    review = report.get('review_summary',{})
    def count(value):
        return len(value) if isinstance(value,list) else (value or 0)
    memo.append(p(f"Numeric source review: {count(review.get('human_confirmed'))} of {count(review.get('required'))} readings have recorded human confirmation. AI or synthetic checks do not count as human confirmation."))
    memo.append(p('Please respond by issue ID, stating whether each item is accepted, needs further evidence, or remains disputed. A focused reinspection can address differing observations at the named locations.'))

    letter = [p(state), p(f'Re Case {cid} roof estimate review'), p('Dear Claim Representative,'), p(opening)]
    for item in computed[:3]:
        raw = gs.get(item['claim_id'], {})
        text = raw.get('requested_action') or raw.get('mapping_rationale') or 'Please review the documented difference.'
        letter.append(p(item['claim_id'] + ': ' + text))
    if totals and totals.get('downward_changes') not in (None,'0.00'):
        letter.append(p(f"The comparison includes {dollars(totals['downward_changes'])} in reductions as well as the documented increases."))
    letter.append(p('The memorandum links each issue to its source evidence and calculation and identifies unresolved coverage questions. These construction figures are not a calculation of the amount currently payable.'))
    letter.append(p('Please provide an itemized response, identify any additional evidence needed, and explain any different measurement, scope, pricing, or policy basis. We would welcome a focused reinspection where observations differ.'))
    letter += [p('Thank you for reviewing the documentation.'), p('Sincerely,'), p(case.get('sender',{}).get('name') or 'Sender to be confirmed'),
               p('Attachments: reconciliation memorandum, operator checklist, evidence appendix. A revised licensed Xactimate estimate should be attached after operator reconciliation.')]

    checklist = [p(state), p('Use the verified record to prepare a separately saved estimate in licensed Xactimate. This worksheet does not create an ESX file.')]
    common = [
        'Confirm active estimate versions, full page counts, source values and the exact pricing location and month.',
        'Reconcile Sketch faces, slopes, lengths and openings against the measurement record; keep plan and surface area distinct.',
        'Inspect each selected item and its labor, material and equipment components in the licensed installation.',
        'Check automatic waste, manual formulas, macros, starter and ridge for overlap.',
        'Attach item notes identifying the location, observation, exhibit, quantity calculation and inclusion check.'
    ]
    checklist += [p('[ ] ' + text) for text in common]
    for item in report.get('groups', []):
        raw = gs.get(item['claim_id'], {})
        checklist.append(p('[ ] ' + item['claim_id'] + ': ' + (raw.get('requested_action') or 'Resolve the named item') + ' Expected direct change: ' + dollars(item.get('delta')) + '. Existing lines: ' + ', '.join(item.get('baseline',[])) + '. Proposed lines: ' + ', '.join(item.get('proposed',[])) + '.'))
        ids = set(item.get('baseline',[]) + item.get('proposed',[]))
        matched = [line for line in case.get('lines',[]) if line.get('id') in ids]
        checklist.append(table(['Side / line','Quantity / unit','Unit price','Line total'],[
            [line['side'] + ' / ' + line['id'],str((line.get('quantity') or {}).get('value') or 'Unknown') + ' ' + line.get('unit',''),
             dollars((line.get('unit_price') or {}).get('value')),dollars((line.get('reported_total') or {}).get('value'))] for line in matched]))
        checklist.append(p('Evidence IDs: ' + ', '.join(item.get('source_ids',[])) + '. Confirm the actual selector, activity and included operations in licensed Xactimate; none is fabricated by this worksheet.'))
    checklist += [p('[ ] ' + text) for text in [
        'Document job-specific price overrides and review tax, general O&P and labor minimum bases without defaults.',
        'Run the configured estimate inspection; resolve material findings or record a justified exception.',
        'Retain a Variation report for checkpoint pricing and a separate carrier-versus-contractor scope crosswalk.',
        'Export a readable revised estimate and the native project where authorized; verify that attachments are visible.',
        'Compare exported direct costs and financial totals to this memo. Record every difference and the operator identity.',
        'Confirm any policy/payment treatment separately and obtain the actual sender\'s approval of the exact document version.'
    ]]

    appendix = [p(state), p('Exhibits support the specified observations only. Original bytes and complete hashes are preserved in the case and run manifests.'),
                table(['Exhibit', 'Document and page', 'What it shows'], [[eid, f"{e.get('doc_id')} p. {e.get('page')}", e.get('description','')] for eid,e in evidence.items()])]
    appendix += [heading('Original document index'), table(['Document', 'Version', 'SHA256 prefix'], [[d['id'],d.get('version',''),d.get('sha256','')[:16]] for d in docs.values()])]
    appendix.append(heading('Numerical source trace'))
    from .engine import review_subjects
    for subject,field in review_subjects(case).items():
        source = field.get('source') or {}
        appendix.append(p(f"{subject}: {field.get('value') if field.get('value') is not None else 'Unknown'}; raw {field.get('raw_text','')}; {source.get('doc_id','Unknown')} p. {source.get('page','?')}; crop {source.get('bbox','not recorded')}; unit {field.get('unit') or '?'}."))
    appendix.append(heading('Evidence excerpts'))
    seen = set()
    for e in evidence.values():
        d = docs.get(e.get('doc_id'), {})
        for page in d.get('pages', []):
            key = (d.get('id'), page.get('page'))
            if page.get('page') == e.get('page') and page.get('image_path') and key not in seen:
                seen.add(key)
                appendix.append({'type':'image','path':page['image_path'],'caption':f"{e['id']} | {d['id']} page {page['page']} | {e.get('description','')}"})
    for d in docs.values():
        if d.get('role') in ('baseline','proposed'):
            for page in d.get('pages',[]):
                key = (d['id'],page.get('page'))
                if page.get('image_path') and key not in seen:
                    seen.add(key)
                    appendix.append({'type':'image','path':page['image_path'],'caption':f"Source estimate | {d['id']} page {page['page']} | Numerical crops are indexed above."})
    if context:
        appendix.append(heading('Context packs considered'))
        appendix.append(p('Context packs guide questions and writing. Their prices and rules are not automatically applied to this claim.'))
        for entry in context.get('applicable', []):
            pack = entry.get('pack',entry) if isinstance(entry,dict) else {}
            appendix.append(p(str(pack.get('id','Context')) + ' ' + str(pack.get('version','')) + ': ' + str(pack.get('title',''))))
        for warning in context.get('warnings', []):
            appendix.append(p(str(warning)))
    return {'reconciliation_memo': {'title':'Roof Estimate Reconciliation Memorandum','blocks':memo},
            'cover_letter': {'title':'Request for Roof Estimate Review','blocks':letter},
            'operator_checklist': {'title':'Xactimate Operator Checklist','blocks':checklist},
            'evidence_appendix': {'title':'Roof Estimate Evidence Appendix','blocks':appendix}}


def markdown(document):
    lines = ['# ' + document['title'], '']
    def cell(value):
        return str(value).replace('|','\\|').replace('\n',' ')
    for block in document['blocks']:
        if block['type'] == 'p':
            lines += [block['text'], '']
        elif block['type'] == 'heading':
            lines += ['## ' + block['text'], '']
        elif block['type'] == 'table':
            lines += ['| ' + ' | '.join(map(cell,block['headers'])) + ' |', '| ' + ' | '.join('---' for _ in block['headers']) + ' |']
            lines += ['| ' + ' | '.join(map(cell,row)) + ' |' for row in block['rows']]
            lines.append('')
        elif block['type'] == 'image':
            lines += ['![Evidence page](' + block['path'] + ')', block['caption'], '']
    return '\n'.join(lines)


def _docx(document, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    for name in ('Title','Heading 1','Heading 2'):
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.color.rgb = RGBColor(0,0,0)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.add_paragraph(document['title'], 'Title')
    for block in document['blocks']:
        typ = block['type']
        if typ == 'p':
            doc.add_paragraph(block['text'])
        elif typ == 'heading':
            doc.add_paragraph(block['text'], 'Heading 1')
        elif typ == 'table':
            t = doc.add_table(rows=1, cols=len(block['headers']))
            t.style = 'Light Shading Accent 1'
            for cell, value in zip(t.rows[0].cells,block['headers']):
                cell.text = str(value)
            repeat = OxmlElement('w:tblHeader')
            t.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in block['rows']:
                cells = t.add_row().cells
                for cell, value in zip(cells,row):
                    cell.text = str(value)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
            doc.add_paragraph()
        elif typ == 'image':
            image = path.parent / block['path']
            if image.exists():
                from PIL import Image
                with Image.open(image) as img:
                    width = min(6.5, 7.5 * img.width / img.height)
                doc.add_picture(str(image),width=Inches(width))
                doc.add_paragraph(block['caption'], 'Caption')
    foot = section.footer.paragraphs[0]
    foot.text = 'Review draft | '
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'),'PAGE')
    foot._p.append(fld)
    doc.core_properties.author = ''
    doc.core_properties.title = document['title']
    doc.save(path)


def _pdf(document, path, fonts_root=None):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font, bold = 'Helvetica', 'Helvetica-Bold'
    if fonts_root and (Path(fonts_root)/'DejaVuSans.ttf').exists():
        for name, file in [('XRRegular','DejaVuSans.ttf'),('XRBold','DejaVuSans-Bold.ttf')]:
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name,str(Path(fonts_root)/file)))
        font,bold='XRRegular','XRBold'
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('XRBody',fontName=font,fontSize=9,leading=12,spaceAfter=5,splitLongWords=True))
    styles.add(ParagraphStyle('XRTitle',fontName=bold,fontSize=20,leading=24,spaceAfter=14))
    styles.add(ParagraphStyle('XRHeading',fontName=bold,fontSize=12,leading=16,spaceBefore=10,spaceAfter=7,keepWithNext=True))
    styles.add(ParagraphStyle('XRCell',fontName=font,fontSize=8,leading=11,splitLongWords=True))
    styles.add(ParagraphStyle('XRCaption',fontName=font,fontSize=8,leading=11,spaceAfter=8))
    para = lambda text, style='XRBody': Paragraph(html.escape(str(text)).replace('\n','<br/>'),styles[style])
    flow = [para(document['title'],'XRTitle')]
    for block in document['blocks']:
        if block['type'] == 'p':
            flow.append(para(block['text']))
        elif block['type'] == 'heading':
            flow.append(para(block['text'],'XRHeading'))
        elif block['type'] == 'table':
            rows = [[para(v,'XRCell') for v in block['headers']]]
            rows += [[para(v,'XRCell') for v in row] for row in block['rows']]
            widths = [468/len(block['headers'])]*len(block['headers'])
            t = Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8edf1')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#84919d')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f7f8')])]))
            flow += [t,Spacer(1,8)]
        elif block['type'] == 'image':
            image = path.parent / block['path']
            if image.exists():
                from PIL import Image as PILImage
                with PILImage.open(image) as img:
                    width = min(468,450*img.width/img.height)
                    height = width*img.height/img.width
                flow.append(KeepTogether([Image(str(image),width=width,height=height),para(block['caption'],'XRCaption')]))
    def page_footer(canvas, doc):
        canvas.setFont(font,8)
        canvas.setFillColor(colors.HexColor('#526170'))
        canvas.drawString(54,30,'Review draft')
        canvas.drawRightString(558,30,str(doc.page))
    SimpleDocTemplate(str(path),pagesize=(612,792),leftMargin=54,rightMargin=54,topMargin=44,bottomMargin=48,title=document['title'],author='').build(flow,onFirstPage=page_footer,onLaterPages=page_footer)


def render_all(case, report, case_root, output, formats=('md',), context=None, assets_root=None):
    import shutil
    output = Path(output)
    output.mkdir(parents=True,exist_ok=True)
    documents = build_documents(case,report,context)
    warnings = []
    for name, document in documents.items():
        for block in document['blocks']:
            if block['type'] != 'image':
                continue
            original = within(case_root,block['path'])
            target = output / 'evidence_pages' / (sha_file(original)[:20]+'.png')
            target.parent.mkdir(exist_ok=True)
            shutil.copy2(original,target)
            block['path'] = target.relative_to(output).as_posix()
        (output/(name+'.md')).write_text(markdown(document),encoding='utf-8')
        for fmt, builder in [('docx',_docx),('pdf',_pdf)]:
            if fmt not in formats:
                continue
            try:
                if fmt == 'pdf':
                    builder(document,output/(name+'.pdf'),Path(assets_root)/'fonts' if assets_root else None)
                else:
                    builder(document,output/(name+'.docx'))
            except ImportError as e:
                warnings.append(fmt + ' unavailable: ' + str(e))
    return {'documents': list(documents), 'warnings': warnings}
